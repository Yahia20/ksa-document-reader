"""Command line.

    python -m docreader scan.jpg                 # print the fields and their status
    python -m docreader scan.jpg --json          # full result as JSON
    python -m docreader folder/ --excel out.xlsx # a whole folder into one spreadsheet
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from .pipeline import DocumentReader

MARK = {"verified": "OK ", "read": "?  ", "review": "!! "}
IMAGES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}


def show(path: Path, res) -> None:
    print(f"\n{path.name}: {res.doc_type}  ({res.type_basis}, {res.seconds}s)")
    for name, f in res.fields.items():
        value = f.value if not isinstance(f.value, list) else f"{len(f.value)} lines"
        print(f"  {MARK[f.status]}{name:18s} {value!s:42.42s} {f.source:10s} {f.note}")
    if res.needs_review:
        print(f"  -> needs a person: {', '.join(res.needs_review)}")
    else:
        print("  -> every field verified")


def to_excel(results: list[tuple[Path, object]], out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    fills = {"verified": "C8E6C9", "read": "FFF59D", "review": "FFCDD2"}
    wb = Workbook()
    ws = wb.active
    ws.title = "documents"
    names = sorted({n for _, r in results for n in r.fields})
    ws.append(["file", "type", "needs review"] + names)
    for path, r in results:
        row = [path.name, r.doc_type, ", ".join(r.needs_review)]
        row += [_cell(r.fields[n].value) if n in r.fields else "" for n in names]
        ws.append(row)
        for i, n in enumerate(names, start=4):
            if n in r.fields:
                ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=fills[r.fields[n].status])
    wb.save(out)


def _cell(v):
    if isinstance(v, list):
        return f"{len(v)} lines"
    return v if isinstance(v, (int, float, str)) or v is None else str(v)


def main() -> None:
    ap = argparse.ArgumentParser(description="Read passports, IDs, visas, ZATCA invoices and boarding passes.")
    ap.add_argument("path", type=Path)
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--excel", type=Path, help="write all results to this .xlsx file")
    a = ap.parse_args()

    files = sorted(p for p in a.path.iterdir() if p.suffix.lower() in IMAGES) if a.path.is_dir() else [a.path]
    reader = DocumentReader()
    results = []
    for f in files:
        res = reader.read(f)
        results.append((f, res))
        if a.json:
            print(json.dumps(res.to_dict(), ensure_ascii=False, indent=2))
        else:
            show(f, res)
    if a.excel:
        to_excel(results, a.excel)
        print(f"\nwrote {a.excel}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
